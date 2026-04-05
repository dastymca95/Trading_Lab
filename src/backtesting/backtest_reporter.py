from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


# ═══════════════════════════════════════════════════════════════════════
# EXCEL
# ═══════════════════════════════════════════════════════════════════════

HFILL = PatternFill('solid', fgColor='1F2937')
HFONT = Font(color='FFFFFF', bold=True)
thin  = Side(style='thin', color='D1D5DB')
BRD   = Border(left=thin, right=thin, top=thin, bottom=thin)


def write_df_sheet(wb, name, df):
    ws = wb.create_sheet(name)
    if df is None or len(df) == 0:
        ws['A1'] = 'Sin datos'
        return ws
    ws.append(df.columns.tolist())
    for row in df.itertuples(index=False):
        ws.append(list(row))
    for cell in ws[1]:
        cell.fill = HFILL
        cell.font = HFONT
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for row in ws.iter_rows(min_row=2, max_row=min(ws.max_row, 500)):
        for cell in row:
            cell.border = BRD
    ws.freeze_panes = 'A2'
    for col in ws.iter_cols(max_row=min(ws.max_row, 200)):
        letter = col[0].column_letter
        width  = max((len(str(cell.value)) for cell in col if cell.value), default=10)
        ws.column_dimensions[letter].width = min(width + 2, 28)
    return ws
